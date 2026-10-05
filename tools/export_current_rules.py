"""지금 앱에 적용된 점수·추세선 규칙을 엑셀로 정리한다."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.signals import (
    DEFAULT_CUTS_CRYPTO,
    DEFAULT_CUTS_STOCK,
    DEFAULT_WEIGHTS,
    SCORE_BASE,
    SCORE_HI,
    SCORE_LO,
    SIGNAL_RULE_VERSION,
    score_to_pct,
    _action_from_pct,
)

OUT = ROOT / "현재_규칙.xlsx"

THIN = Border(
    left=Side(style="thin", color="E5E7EB"),
    right=Side(style="thin", color="E5E7EB"),
    top=Side(style="thin", color="E5E7EB"),
    bottom=Side(style="thin", color="E5E7EB"),
)
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=12)
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="F8CBAD")
GRAY = PatternFill("solid", fgColor="F3F4F6")
BLUE = PatternFill("solid", fgColor="DDEBF7")
WRAP = Alignment(vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")
TITLE = Font(size=18, bold=True, color="1F4E79")
SECTION = Font(bold=True, color="1F4E79", size=13)


def paint_header(ws, n_cols: int) -> None:
    ws.row_dimensions[1].height = 24
    for col in range(1, n_cols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN
    ws.freeze_panes = "A2"


def paint_body(ws, score_col: int | None = None) -> None:
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for i, cell in enumerate(row, 1):
            cell.alignment = WRAP
            cell.border = THIN
            if score_col and i == score_col:
                cell.alignment = CENTER
                cell.font = Font(bold=True, size=12)
                val = cell.value
                if isinstance(val, (int, float)):
                    if val > 0:
                        cell.fill = GREEN
                        cell.value = f"+{int(val)}"
                    elif val < 0:
                        cell.fill = RED
                    else:
                        cell.fill = GRAY
                elif isinstance(val, str) and val.startswith("+"):
                    cell.fill = GREEN
                elif isinstance(val, str) and val.startswith("-"):
                    cell.fill = RED


def add_sheet(wb, title, headers, rows, widths, score_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    paint_header(ws, len(headers))
    paint_body(ws, score_col)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = 36
    return ws


def main() -> None:
    wb = Workbook()
    how = wb.active
    how.title = "이렇게 봐요"
    how["A1"] = f"지금 적용 규칙 · 앱 v{SIGNAL_RULE_VERSION}"
    how["A1"].font = TITLE
    how.merge_cells("A1:B1")
    lines = [
        "",
        "이 파일은 앱이 실제로 쓰는 점수·추세선·컷을 정리한 것입니다.",
        "",
        "순서",
        "1. 「기본 점수」 항목을 해당되면 더하거나 뺍니다. 기본은 항상 10점입니다.",
        f"2. 합산 점수를 %로 바꿉니다. {SCORE_LO}점 이하=0%, {SCORE_BASE}점=50%, {SCORE_HI}점 이상=100%.",
        "3. %로 매수 / 매도 / 홀딩을 정합니다. 주식과 코인 컷은 같습니다.",
        "4. 홀딩이면 끝입니다. 매수나 매도면 미국 주식·오늘 조회만 옵션 점수를 더합니다.",
        "",
        "가까운 가격",
        "하루 변동폭(ATR)의 약 20%, 또는 주가의 0.6% 중 더 큰 값.",
        "",
        "시트 안내",
        "기본 점수 — 차트·지표로 매기는 점수",
        "조회기간 추세 — 3개월·6개월·1년만. 상승선 기울기가 아니라 스윙+이평으로 판정",
        "옵션 점수 — 미국 주식, 매수/매도일 때 추가",
        "매수 매도 기준 — 합산 % 컷",
        "점수별 % — 점수마다 앱이 매기는 %",
        "추세선 긋는 법 — 장기·단기 상승선, 하락선",
        "추세 판정 — 조회기간 추세 항목이 상승/하락/횡보를 나누는 방법",
    ]
    for i, line in enumerate(lines, 2):
        how[f"A{i}"] = line
        if line in ("순서", "가까운 가격", "시트 안내"):
            how[f"A{i}"].font = SECTION
        how[f"A{i}"].alignment = WRAP
    how.column_dimensions["A"].width = 88
    how.row_dimensions[1].height = 28

    add_sheet(
        wb,
        "기본 점수",
        ["항목", "이럴 때", "점수"],
        [
            ["기본", "시작할 때 항상 줌", 10],
            ["조회기간 추세", "(3개월,6개월,1년 조회에만 적용) 조회기간 상승 · 1개월(1시간봉) 상승", -1],
            ["조회기간 추세", "(3개월,6개월,1년 조회에만 적용) 조회기간 상승 · 1개월(1시간봉) 하락", 0],
            ["조회기간 추세", "(3개월,6개월,1년 조회에만 적용) 조회기간 하락 · 1개월(1시간봉) 상승", 0],
            ["조회기간 추세", "(3개월,6개월,1년 조회에만 적용) 조회기간 하락 · 1개월(1시간봉) 하락", 1],
            ["조회기간 추세", "(3개월,6개월,1년 조회에만 적용) 조회기간 횡보 · 1개월(1시간봉) 하락", 1],
            ["조회기간 추세", "(3개월,6개월,1년 조회에만 적용) 조회기간 횡보 · 1개월(1시간봉) 상승", -1],
            ["조회기간 추세", "1개월 조회, 또는 1개월(1시간봉)이 횡보/없음", 0],
            ["하락 추세선 근접", "하락 추세선에 근접 했을 때, 돌파하면 무효", -1],
            ["상승 추세선 근접", "상승 추세선 근접 했을 때, 이탈하면 무효", 1],
            ["지지 근접", "지지선 근접이고 강도 4 이상일 때", 1],
            ["저항 근접", "저항선 바로 옆이고, 강도 4 이상일 때", -1],
            ["최대 매물 (POC)", "현재가가 거래가 가장 많았던 가격 근접 했을 때, 이탈 시 무효", 1],
            ["밸류 하단 (VAL)", "현재가가 싼 구간 아래일때 (추세 무관)", 1],
            ["밸류 상단 (VAH)", "현재가가 비싼 구간 위일때 (거리 무관)", -1],
            ["RSI", "35 이하 (너무 많이 떨어짐)", 1],
            ["RSI", "70 이상 (너무 많이 오름)", -1],
            ["20일선", "현재가가 20일선 근처일때, 20일선 완전 이탈시 무효", 1],
            ["60일선", "현재가가 60일선 근처일때, 60일선 완전 이탈시 무효", 1],
            ["180일선", "현재가가 장기 이평 근처일 때. 6개월 조회는 180일선, 1년 조회는 300일선, 완전이탈시 무효", 1],
            ["20일선 60일선 교차", "20일선 방향이 하방으로 떨어지면서 60일선 아래로 떨어지기 시작할때, 떨어지고 4봉이상 지나면 무효", -1],
            ["1개월 하락률", "한 달동안 15% 이상 25%미만 하락 했을 때", 1],
            ["1개월 하락률", "한 달 동안 25% 이상 35% 미만 떨어짐", 2],
            ["1개월 하락률", "한 달 동안 35% 이상 떨어짐", 3],
            ["6개월 상승률", "6개월 동안 800% 이상 오름, 횡보 추세나 하락 추세시 무효", -3],
            ["6개월 상승률", "6개월 동안 600% 이상 오름, 횡보 추세나 하락 추세시 무효", -2],
            ["6개월 상승률", "6개월 동안 300% 이상 600% 미만 오름, 횡보 추세나 하락 추세시 무효", -1],
            ["단기 급상승", "1개 봉만에 20% 이상 상승했을 시", -1],
            ["신고가 달성", "6개월 상승률이 800% 미만이면서 위에 저항이 없는 신고가의 경우, 3봉 이후 무효", 1],
            ["신고가 저항", "6개월 상승률이 800% 이상이면서 위에 저항이 없는 신고가의 경우, 2봉 이후 무효", -1],
            ["신고가 이탈", "신고가 돌파 후 5봉 이내에 다시 하락 추세선 안으로 현재가가 내려왔을 때, 3봉 이후 무효", -1],
        ],
        [18, 78, 10],
        score_col=3,
    )

    add_sheet(
        wb,
        "조회기간 추세",
        ["조회기간 추세", "1개월(1시간봉) 추세", "점수", "비고"],
        [
            ["상승", "상승", -1, "3개월·6개월·1년만"],
            ["상승", "하락", 0, ""],
            ["하락", "상승", 0, ""],
            ["하락", "하락", 1, ""],
            ["횡보", "상승", -1, ""],
            ["횡보", "하락", 1, ""],
            ["상승/하락/횡보", "횡보 또는 없음", 0, ""],
            ["(1개월·2개월 조회)", "해당 없음", 0, "이 항목은 3개월 이상만"],
        ],
        [22, 22, 10, 36],
        score_col=3,
    )

    add_sheet(
        wb,
        "옵션 점수",
        ["항목", "이럴 때", "점수"],
        [
            ["옵션", "기존 결과가 홀딩이면 옵션은 보지 않음", 0],
            ["옵션", "기존이 매도인데, 만기가 14일 안이고, 위쪽에 콜 벽이 두껍고 아래 풋 벽이 얇음", -1],
            ["옵션", "기존이 매도인데, 반대로 아래 풋 벽이 두껍고 위 콜 벽이 얇음", 1],
            ["옵션", "기존이 매도인데, 벽이 멀거나 둘 다 얇음", 0],
            ["옵션", "기존이 매수인데, 아래 풋 벽이 얇고 위 콜 벽이 두꺼움", -1],
            ["옵션", "기존이 매수인데, 위와 다르면", 0],
        ],
        [12, 78, 10],
        score_col=3,
    )
    opt = wb["옵션 점수"]
    opt["A9"] = "참고: 미국 주식, 오늘 조회만. 근처는 현재가에서 5% 안, 멀면 8% 밖."
    opt["A9"].font = Font(italic=True, color="666666")
    opt.merge_cells("A9:C9")

    cuts = wb.create_sheet("매수 매도 기준")
    cuts.merge_cells("A1:B1")
    cuts.merge_cells("D1:E1")
    cuts["A1"] = "코인"
    cuts["D1"] = "주식"
    for cell in (cuts["A1"], cuts["D1"]):
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    cuts["A2"] = "합산 %"
    cuts["B2"] = "제안"
    cuts["D2"] = "합산 %"
    cuts["E2"] = "제안"
    for col in (1, 2, 4, 5):
        cell = cuts.cell(2, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    rows = [
        ["73% 이상", "강한 매수"],
        ["65% 이상 ~ 73% 미만", "매수"],
        ["55% 이상 ~ 65% 미만", "약한 매수"],
        ["35% 초과 ~ 55% 미만", "홀딩"],
        ["25% 초과 ~ 35% 이하", "약한 매도"],
        ["20% 초과 ~ 25% 이하", "매도"],
        ["20% 이하", "강한 매도"],
    ]
    fills = (GREEN, GREEN, PatternFill("solid", fgColor="E2EFDA"), YELLOW, RED, RED, PatternFill("solid", fgColor="F4B183"))
    for i, ((pct, act), fill) in enumerate(zip(rows, fills), 3):
        cuts.cell(i, 1, pct)
        cuts.cell(i, 2, act)
        cuts.cell(i, 4, pct)
        cuts.cell(i, 5, act)
        for col in (1, 2, 4, 5):
            cell = cuts.cell(i, col)
            cell.fill = fill
            cell.border = THIN
            cell.alignment = CENTER
            if col in (2, 5):
                cell.font = Font(bold=True, size=12)
        cuts.row_dimensions[i].height = 28
    cuts["A11"] = (
        f"점수를 %로 바꾸는 법: {SCORE_LO}점 이하=0%, {SCORE_BASE}점(기본)=50%, {SCORE_HI}점 이상=100%."
    )
    cuts["A11"].font = Font(italic=True, color="666666")
    cuts.merge_cells("A11:E11")
    cuts["A12"] = "가까운 가격: 하루 변동폭의 약 20%, 또는 주가의 0.6% 중 더 큰 값."
    cuts["A12"].font = Font(italic=True, color="666666")
    cuts.merge_cells("A12:E12")
    assert DEFAULT_CUTS_STOCK == DEFAULT_CUTS_CRYPTO
    cuts.column_dimensions["A"].width = 28
    cuts.column_dimensions["B"].width = 14
    cuts.column_dimensions["C"].width = 4
    cuts.column_dimensions["D"].width = 28
    cuts.column_dimensions["E"].width = 14

    pct_rows = []
    lo_span = SCORE_BASE - SCORE_LO
    hi_span = SCORE_HI - SCORE_BASE
    for s in range(-3, SCORE_HI + 4):
        pct = score_to_pct(s)
        action = _action_from_pct(pct, DEFAULT_CUTS_STOCK)
        if s < SCORE_LO:
            process = f"{s}점 → {SCORE_LO}점으로 보고 0%"
        elif s > SCORE_HI:
            process = f"{s}점 → {SCORE_HI}점으로 보고 100%"
        elif s <= SCORE_BASE:
            process = f"({s}-{SCORE_LO})/{lo_span}×50 = {(s - SCORE_LO) / lo_span * 50:.3f} → {pct}%"
        else:
            process = (
                f"50+({s}-{SCORE_BASE})/{hi_span}×50 = "
                f"{50 + (s - SCORE_BASE) / hi_span * 50:.3f} → {pct}%"
            )
        formula = (
            f"=ROUND(IF(A{len(pct_rows)+2}<={SCORE_LO},0,IF(A{len(pct_rows)+2}>={SCORE_HI},100,"
            f"IF(A{len(pct_rows)+2}<={SCORE_BASE},(A{len(pct_rows)+2}-{SCORE_LO})/{lo_span}*50,"
            f"50+(A{len(pct_rows)+2}-{SCORE_BASE})/{hi_span}*50))),0)"
        )
        pct_rows.append([s, pct, formula, process, action, action])

    ws_pct = add_sheet(
        wb,
        "점수별 %",
        ["점수", "앱 평가 %", "엑셀 수식 %", "계산 과정", "주식 제안", "코인 제안"],
        [[r[0], r[1], r[2], r[3], r[4], r[5]] for r in pct_rows],
        [10, 14, 22, 52, 14, 14],
        score_col=None,
    )
    for r in range(2, ws_pct.max_row + 1):
        ws_pct.cell(r, 3).value = pct_rows[r - 2][2]
        for col in range(1, 7):
            ws_pct.cell(r, col).alignment = WRAP if col == 4 else CENTER
            ws_pct.cell(r, col).border = THIN
        ws_pct.row_dimensions[r].height = 22
    ws_pct["A1"].fill = HEADER_FILL
    note = wb.create_sheet("점수별 % 계산식")
    note["A1"] = "점수 → % 환산"
    note["A1"].font = TITLE
    note.merge_cells("A1:B1")
    note_lines = [
        "",
        "앱이 쓰는 기준점",
        f"  하한 LO = {SCORE_LO}점  →  0%",
        f"  기본 BASE = {SCORE_BASE}점  →  50%",
        f"  상한 HI = {SCORE_HI}점  →  100%",
        "",
        "계산식 (점수를 s라고 할 때)",
        f"  s ≤ {SCORE_LO}  →  0%",
        f"  {SCORE_LO} < s ≤ {SCORE_BASE}  →  ROUND( (s-{SCORE_LO})/{lo_span} × 50 )",
        f"  {SCORE_BASE} < s < {SCORE_HI}  →  ROUND( 50 + (s-{SCORE_BASE})/{hi_span} × 50 )",
        f"  s ≥ {SCORE_HI}  →  100%",
        "",
        "엑셀 수식 예 (A열에 점수)",
        f"  =ROUND(IF(A2<={SCORE_LO},0,IF(A2>={SCORE_HI},100,IF(A2<={SCORE_BASE},(A2-{SCORE_LO})/{lo_span}*50,50+(A2-{SCORE_BASE})/{hi_span}*50))),0)",
        "",
        "앱은 Python round(은행가 반올림)를 씁니다. 0.5가 정확히 떨어지면 짝수로 갑니다.",
        f"기본 배점 {DEFAULT_WEIGHTS['base']}점은 {score_to_pct(int(DEFAULT_WEIGHTS['base']))}%입니다.",
    ]
    for i, line in enumerate(note_lines, 2):
        note[f"A{i}"] = line
        note[f"A{i}"].alignment = WRAP
        if line in ("앱이 쓰는 기준점", "계산식 (점수를 s라고 할 때)", "엑셀 수식 예 (A열에 점수)"):
            note[f"A{i}"].font = SECTION
    note.column_dimensions["A"].width = 96

    line_ws = wb.create_sheet("추세선 긋는 법")
    line_ws["A1"] = "상승·하락 추세선"
    line_ws["A1"].font = TITLE
    line_ws.merge_cells("A1:B1")
    line_lines = [
        "",
        "공통",
        "스윙 고점·저점은 좌우 4봉. 조회기간으로 자른 차트 안의 스윙만 씁니다.",
        "주식은 1개월 1시간봉, 2개월 4시간봉, 3개월·6개월·1년 일봉입니다. 코인은 1개월 12시간봉, 2·3개월·6개월·1년은 일봉입니다.",
        "",
        "하락 추세선",
        "마지막 스윙 고점 두 개를 잇고, 마지막 봉까지 연장합니다. 기울기 제한 없음.",
        "",
        "장기 상승 추세선 (1개월·2개월·3개월·6개월·1년 모두 같음)",
        "조회기간 전체 스윙 저점을 씁니다.",
        "1. 가장 낮은 점(같은 가격이면 가장 오래된 점)과 가장 최근 저점을 잇습니다.",
        "2. 그 직선 아래 스윙 저점이 있으면, 그중 가장 최근 점을 새 오른쪽 끝으로 둡니다.",
        "3. 아래에 저점이 없을 때까지 2를 반복합니다.",
        "4. 최근 스윙이 유일한 최저이면 마지막 두 저점을 긋습니다. 이때만 하방이 됩니다.",
        "점수 항목 「상승 추세선 근접」은 이 장기선을 봅니다.",
        "",
        "단기 상승 추세선",
        "같은 방식이지만 최근 스윙 저점 5개만 씁니다. 차트에 점선으로 추가로 그립니다.",
        "",
        "가까운 가격(근접 판정)",
        "max(ATR × 0.20, 가격 × 0.6%). 상승선은 선 위·근처, 하락선은 선 아래·근처일 때 근접입니다.",
    ]
    for i, line in enumerate(line_lines, 2):
        line_ws[f"A{i}"] = line
        line_ws[f"A{i}"].alignment = WRAP
        if line in (
            "공통",
            "하락 추세선",
            "장기 상승 추세선 (1개월·2개월·3개월·6개월·1년 모두 같음)",
            "단기 상승 추세선",
            "가까운 가격(근접 판정)",
        ):
            line_ws[f"A{i}"].font = SECTION
    line_ws.column_dimensions["A"].width = 96

    trend_ws = wb.create_sheet("추세 판정")
    trend_ws["A1"] = "조회기간 추세 항목의 상승/하락/횡보"
    trend_ws["A1"].font = TITLE
    trend_ws.merge_cells("A1:B1")
    trend_lines = [
        "",
        "이 판정은 상승추세선 기울기를 보지 않습니다. 스윙 구조와 이평만 봅니다.",
        "조회기간 쪽은 그 조회 차트(an.trend), 1개월 쪽은 1시간봉 30일에 같은 방법을 씁니다.",
        "",
        "1. 스윙(좌우 4봉)의 마지막 두 저점·두 고점",
        "   저점·고점이 둘 다 높아지면 상승, 둘 다 낮아지면 하락.",
        "2. 20일선·60일선과 현재가",
        "   20일선이 60일선보다 0.5% 이상 위이고 가격이 20일선 위면 상승.",
        "   20일선이 60일선보다 0.5% 이상 아래이고 가격이 20일선 아래면 하락.",
        "3. 스윙 구조가 나오면 그걸 쓰고, 스윙이 횡보면 이평을 씁니다.",
        "   60일선이 아직 없으면(짧은 봉) 스윙+20일선, 또는 기간 첫·끝 종가로 보조합니다.",
        "",
        "그래서 차트 상승선이 위를 향해도 조회기간 추세는 하락일 수 있습니다.",
    ]
    for i, line in enumerate(trend_lines, 2):
        trend_ws[f"A{i}"] = line
        trend_ws[f"A{i}"].alignment = WRAP
        if line.startswith("1.") or line.startswith("2.") or line.startswith("3."):
            trend_ws[f"A{i}"].font = Font(bold=True)
    trend_ws.column_dimensions["A"].width = 96

    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
